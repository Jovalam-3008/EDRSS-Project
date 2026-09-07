"""Construye el informe académico EDRSS V2 a partir de la plantilla institucional."""

from __future__ import annotations

from pathlib import Path

try:
    import matplotlib.pyplot as plt
    import numpy as np
    import pandas as pd
except ModuleNotFoundError:
    plt = np = pd = None
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT.parent / "PROYECTO PRODUCTIVO III" / "deliverables" / "Informe_Proyecto_Productivo_Cobranza_Temprana.docx"
OUTPUT_DIR = ROOT / "deliverables"
OUTPUT = OUTPUT_DIR / "Informe_EDRSS_V2.docx"
FIGURE_DIR = ROOT / ".codex_tmp" / "report_figures"

BLUE = "17365D"
MID_BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_repeat_table_header(row) -> None:
    repeat_header(row)


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(text))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
        if widths:
            cell.width = widths[index]
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if row_number % 2:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
            if widths:
                cell.width = widths[index]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        run.bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.add_run(item)


def add_heading(doc, text: str, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_caption(doc, text: str):
    paragraph = doc.add_paragraph(text, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_picture(doc, path: Path, caption: str, width=6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_field(paragraph, instruction: str, placeholder: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction_node, separate, text, end):
        run._r.append(node)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].paragraph_format.space_before = Pt(15)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(11)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.bold = True
    styles["Caption"].font.color.rgb = RGBColor(89, 89, 89)


def configure_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Proyecto Productivo  EDRSS V2     ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)
        add_field(paragraph, "PAGE", "1")


def create_figures() -> dict[str, Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    ready = {
        "models": FIGURE_DIR / "model_comparison.png",
        "capacity": FIGURE_DIR / "capacity.png",
        "features": FIGURE_DIR / "feature_importance.png",
    }
    if all(path.exists() for path in ready.values()):
        return ready
    if plt is None or np is None or pd is None:
        raise RuntimeError("Las figuras deben generarse antes de construir el DOCX.")
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})

    model_path = FIGURE_DIR / "model_comparison.png"
    models = ["Regresión logística", "HistGradientBoosting", "Random Forest"]
    auc = [0.817919, 0.864995, 0.867025]
    fig, ax = plt.subplots(figsize=(8.4, 3.5))
    colors = ["#A6A6A6", "#5B9BD5", "#17365D"]
    bars = ax.barh(models, auc, color=colors)
    ax.set_xlim(0.75, 0.89)
    ax.set_xlabel("ROC AUC medio en cinco folds temporales")
    ax.grid(axis="x", color="#D9E2F3", linewidth=0.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    for bar, value in zip(bars, auc):
        ax.text(value + 0.002, bar.get_y() + bar.get_height()/2, f"{value:.3f}", va="center", fontweight="bold")
    fig.tight_layout()
    fig.savefig(model_path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    capacity_path = FIGURE_DIR / "capacity.png"
    levels = [5, 10, 20]
    precision = [77.74, 69.54, 58.79]
    recall = [18.82, 33.67, 56.92]
    x = np.arange(len(levels))
    width = 0.34
    fig, ax = plt.subplots(figsize=(8.4, 3.8))
    ax.bar(x - width/2, recall, width, label="Riesgo capturado", color="#4472C4")
    ax.bar(x + width/2, precision, width, label="Precisión", color="#5BC0EB")
    ax.set_xticks(x, [f"{value}%" for value in levels])
    ax.set_xlabel("Capacidad de contacto")
    ax.set_ylabel("Porcentaje")
    ax.set_ylim(0, 90)
    ax.grid(axis="y", color="#E7E6E6", linewidth=0.8)
    ax.legend(frameon=False, loc="upper center", ncol=2)
    ax.spines[["top", "right"]].set_visible(False)
    for container in ax.containers:
        ax.bar_label(container, fmt="%.1f", padding=3, fontsize=9)
    fig.tight_layout()
    fig.savefig(capacity_path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    feature_path = FIGURE_DIR / "feature_importance.png"
    source = ROOT.parent / "PROYECTO PRODUCTIVO III" / "outputs" / "permutation_importance.csv"
    importance = pd.read_csv(source).head(8).sort_values("importance_mean")
    labels = {
        "log1p_sumamnt_ma_rech90": "Monto recargas 90d log",
        "sumamnt_ma_rech90": "Monto recargas 90d",
        "log1p_cnt_ma_rech90": "Conteo recargas 90d log",
        "recharge_recency_vs_tenure": "Recencia respecto a antigüedad",
        "sumamnt_ma_rech30": "Monto recargas 30d",
        "last_rech_date_ma": "Días desde última recarga",
        "medianmarechprebal90": "Saldo previo mediano 90d",
        "cnt_ma_rech90": "Conteo recargas 90d",
    }
    fig, ax = plt.subplots(figsize=(8.4, 4.2))
    ax.barh([labels.get(x, x) for x in importance["variable"]], importance["importance_mean"], color="#4472C4")
    ax.set_xlabel("Disminución de ROC AUC al permutar la variable")
    ax.grid(axis="x", color="#E7E6E6", linewidth=0.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(feature_path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"models": model_path, "capacity": capacity_path, "features": feature_path}


def build_report() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    figures = create_figures()
    doc = Document(str(REFERENCE))
    clear_body(doc)
    configure_styles(doc)
    configure_footer(doc)

    # Portada
    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Early Delinquency Risk Scoring System EDRSS")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("Informe del Proyecto Productivo V2")
    run.bold = True
    run.font.size = Pt(16)
    add_body(doc, "Sistema de scoring para priorizar la cobranza temprana de microcréditos móviles mediante arquitectura Medallion, validación temporal y un dashboard multipágina.")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    authors = [
        "José Luis Valderrama Ampuero", "Edgar Leyvis Villalobos Moreto",
        "Heedeygeer Dante Oropeza Cóndor", "Fernando Jose Castillo Ticona",
        "Daniel Joel Ticona Ccarita", "Pedro Luis Aguayo Trujillo",
    ]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Elaborado por\n").bold = True
    p.add_run("\n".join(authors))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.add_run("Solicitado por\n").bold = True
    p.add_run("Sergio Orizano Salvador")
    p = doc.add_paragraph("Perú 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    doc.add_page_break()

    # Índice
    add_heading(doc, "Índice", 1)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-1" \\h \\z \\u', "El índice se actualizará al abrir el documento en Word")
    doc.add_page_break()

    add_heading(doc, "Introducción", 1)
    add_body(doc, "EDRSS V2 organiza 209,593 operaciones históricas de microcrédito móvil y las convierte en una cola de cobranza temprana. La solución ya funciona sobre Supabase con capas Bronze, Silver y Gold, un modelo Random Forest registrado y un dashboard multipágina que consume únicamente información certificada. El modelo alcanzó un ROC AUC de 0.8435 en un periodo fuera de tiempo. Al contactar el 10 por ciento de mayor riesgo, captura 33.67 por ciento de los fallos observados con una precisión de 69.54 por ciento y un lift de 3.37 frente a una selección aleatoria.")
    add_body(doc, "La V2 responde a una necesidad de orden y mantenibilidad. Los scripts SQL concentran las transformaciones, Python gestiona la conexión, la ingesta y el modelo, y ocho notebooks documentan la ejecución secuencial. El dashboard no vuelve a limpiar ni recalcular la información. Cada hoja consulta una vista Gold o API preparada para esa decisión.")
    add_body(doc, "El resultado debe interpretarse como un sistema de apoyo para ordenar gestiones. No autoriza suspensiones automáticas, sanciones ni decisiones adversas sin revisión humana. La población posterior al 23 de julio de 2016 permanece sin TARGET maduro y se utiliza únicamente para scoring operativo.")

    add_heading(doc, "1 Idea del proyecto", 1)
    add_body(doc, "El proyecto asigna una probabilidad de mora a cada operación de microcrédito móvil a partir de señales conocidas antes de la gestión. Entre estas señales se encuentran la antigüedad de la línea, la recencia y frecuencia de recargas, los montos de préstamos, la actividad de consumo y relaciones entre ventanas de 30 y 90 días.")
    add_body(doc, "El equipo de cobranza utiliza el score para ordenar su cola y seleccionar el número de casos compatible con su capacidad diaria. La política recomendada trabaja con el top K de mayor riesgo, porque un porcentaje de capacidad resulta más estable que un umbral fijo cuando cambia la distribución de los clientes.")

    add_heading(doc, "2 Identificación del problema", 1)
    add_body(doc, "La operación no puede contactar a todos los clientes con igual intensidad. Sin un criterio cuantitativo, una parte de la capacidad puede dirigirse a casos que habrían pagado mientras operaciones que terminarán en mora reciben atención tardía. El problema analítico consiste en ordenar las operaciones para concentrar la mayor cantidad posible de fallos de repago al inicio de la cola.")
    add_body(doc, "El dataset fuente contiene 209,593 filas y 35 variables. La etiqueta original usa 1 para repago exitoso y 0 para fallo. Para que la clase positiva represente el evento que cobra interés, EDRSS define riesgo_mora igual a 1 cuando label es 0. Desde el 24 de julio de 2016 desaparece la clase de riesgo, por lo que esas 58,825 operaciones no se usan para evaluar aprendizaje supervisado.")
    add_table(doc, ["Elemento", "Definición V2"], [
        ["Unidad analítica", "Operación de microcrédito móvil"],
        ["TARGET", "riesgo_mora igual a 1 cuando label es 0"],
        ["Horizonte", "Fallo de repago dentro de cinco días"],
        ["Decisión", "Orden de contacto en cobranza temprana"],
        ["Métrica principal", "ROC AUC fuera de tiempo"],
        ["Política operativa", "Top K según capacidad diaria"],
    ], widths=[Cm(5), Cm(11)])

    add_heading(doc, "3 Objetivos SMART", 1)
    add_body(doc, "El objetivo general consiste en desarrollar y operar un sistema reproducible que ordene operaciones según su riesgo de mora temprana, mida su desempeño fuera de tiempo y publique una cola compatible con la capacidad del equipo de cobranza.")
    add_table(doc, ["Criterio", "Objetivo definido"], [
        ["Específico", "Implementar EDRSS con arquitectura Medallion, scoring y dashboard multipágina"],
        ["Medible", "Conciliar 209,593 filas entre Bronze, Silver y Gold y evaluar ROC AUC, precisión, recall y lift"],
        ["Alcanzable", "Usar Supabase, SQL secuencial, Python y notebooks reproducibles"],
        ["Relevante", "Concentrar la gestión en operaciones con mayor riesgo de fallo"],
        ["Temporal", "Validar con cinco folds y un Test OOT del 13 al 23 de julio de 2016"],
    ], widths=[Cm(3.2), Cm(12.8)])
    add_bullets(doc, [
        "Auditar tipos, nulos, anomalías y ruptura temporal del TARGET.",
        "Construir variables interpretables sin incorporar información posterior a la decisión.",
        "Comparar tres candidatos mediante el mismo protocolo temporal.",
        "Publicar métricas, predicciones y una cola de scoring trazable.",
        "Presentar resultados mediante un tablero orientado a decisiones de cobranza.",
    ])

    add_heading(doc, "4 Búsqueda de literatura académica", 1)
    add_body(doc, "Björkegren y Grissen demostraron que patrones derivados del uso de telefonía móvil pueden predecir el repago y conservar utilidad en evaluaciones posteriores. Esta evidencia respalda el uso de recargas, actividad y recencia como señales de riesgo [1].")
    add_body(doc, "Berg y colaboradores documentaron que las huellas digitales pueden complementar fuentes tradicionales en credit scoring. Su trabajo también refuerza la necesidad de controlar privacidad, trazabilidad y uso responsable del resultado [2].")
    add_body(doc, "Lessmann y colaboradores compararon clasificadores de credit scoring bajo un protocolo común. EDRSS adopta ese principio al evaluar todos los candidatos con los mismos folds y al considerar desempeño promedio, peor periodo y estabilidad [3]. CRISP DM aporta la secuencia de negocio, datos, preparación, modelado, evaluación e implementación [4]. Random Forest se fundamenta en el trabajo de Breiman sobre ensambles de árboles [5].")
    add_body(doc, "La arquitectura Medallion complementa CRISP DM al separar el dato fuente, la información limpia y los productos de consumo. En EDRSS, esta separación mejora la trazabilidad del TARGET, el control de versiones y la conciliación del dashboard [6].")

    add_heading(doc, "5 Justificación e importancia del proyecto", 1)
    add_body(doc, "La justificación de negocio proviene de la capacidad limitada de cobranza. En el Test OOT, 3,161 contactos dirigidos al 10 por ciento de mayor score habrían identificado 2,198 operaciones con mora. Una selección aleatoria necesitaría una capacidad considerablemente mayor para capturar el mismo volumen de riesgo.")
    add_body(doc, "La importancia técnica está en la reproducibilidad. La V2 reemplaza la lógica dispersa por nueve scripts SQL ordenados, módulos Python con responsabilidades reducidas y notebooks que se ejecutan del 00 al 07. Las capas conciliaron exactamente 209,593 registros y el escaneo local confirmó que no existe una contraseña guardada en los archivos.")
    add_body(doc, "La importancia de gobierno radica en separar evaluación y operación. OOF y OOT conservan evidencia de validación. La hoja de priorización muestra únicamente las 58,825 operaciones de SCORING, evitando mezclar datos históricos de evaluación con la cola pendiente.")

    add_heading(doc, "6 Límites y alcances del proyecto", 1)
    add_heading(doc, "6 1 Alcances", 2)
    add_bullets(doc, [
        "Ingesta trazable del CSV y registro de lote en Supabase.",
        "Tipificación, limpieza, reglas semánticas y cuarentena del TARGET.",
        "Cincuenta y ocho variables disponibles en el contrato Gold del modelo.",
        "Validación con cinco folds temporales y embargo de cinco días.",
        "Comparación de regresión logística, Random Forest e HistGradientBoosting.",
        "Publicación de 164,034 predicciones para auditoría y 58,825 casos de scoring para cobranza.",
        "Dashboard de seis hojas con filtros y descarga de la cola priorizada.",
    ])
    add_heading(doc, "6 2 Límites", 2)
    add_bullets(doc, [
        "Los datos son históricos, anonimizados y no incluyen un identificador de cliente que permita controlar recurrencia entre operaciones.",
        "La clase de riesgo desaparece después del 23 de julio de 2016. Esa población no permite medir desempeño supervisado.",
        "No se dispone de costo por contacto, pérdida por mora ni recuperación monetaria. El proyecto no estima ROI financiero.",
        "El modelo identifica asociaciones predictivas y no relaciones causales.",
        "No existen atributos protegidos suficientes para una auditoría completa de equidad.",
        "La operación requiere revisión humana, monitoreo de drift y actualización cuando maduren nuevas etiquetas.",
    ])

    add_heading(doc, "7 Desarrollo del marco teórico y conceptual", 1)
    add_table(doc, ["Concepto", "Aplicación en EDRSS V2"], [
        ["Mora temprana", "Fallo de repago observado dentro del horizonte de cinco días"],
        ["Clasificación binaria", "Estimación de la probabilidad de riesgo_mora"],
        ["Validación temporal", "Entrenamiento con pasado y evaluación en periodos posteriores"],
        ["Embargo", "Separación de cinco días que respeta el horizonte del TARGET"],
        ["Leakage", "Uso indebido de datos futuros o derivados del resultado"],
        ["Medallion", "Progresión desde fuente cruda hacia consumibles certificados"],
        ["Lift", "Concentración de morosos respecto a una selección aleatoria"],
        ["Scoring", "Probabilidad usada para ordenar la cola de cobranza"],
    ], widths=[Cm(4), Cm(12)])
    add_body(doc, "ROC AUC mide la capacidad del modelo para ordenar un caso con mora por encima de uno sin mora. Average Precision resulta útil frente al desbalance porque resume la relación entre precisión y recall. Brier Score mide el error cuadrático de las probabilidades y ayuda a vigilar calibración. Ninguna métrica aislada decide la política de contacto, por lo que el proyecto complementa la evaluación con capacidad del 5, 10 y 20 por ciento.")

    add_heading(doc, "8 Análisis FODA visión misión estrategias y actividades", 1)
    add_heading(doc, "8 1 Análisis FODA", 2)
    add_table(doc, ["Dimensión", "Análisis"], [
        ["Fortalezas", "Arquitectura reproducible, validación temporal, datos conciliados y modelo activo"],
        ["Oportunidades", "Integrar resultados de contacto, estimar retorno y automatizar actualizaciones controladas"],
        ["Debilidades", "Sin llave de cliente, sin costos operativos y con calibración OOT perfectible"],
        ["Amenazas", "Drift, ruptura del TARGET, contactos indebidos y uso punitivo del score"],
    ], widths=[Cm(3.2), Cm(12.8)])
    add_heading(doc, "8 2 Visión", 2)
    add_body(doc, "Contar con una gestión de cobranza temprana basada en datos, trazable y supervisada, que asigne la capacidad de contacto a las operaciones con mayor riesgo sin reemplazar la revisión humana.")
    add_heading(doc, "8 3 Misión", 2)
    add_body(doc, "Transformar señales históricas de recarga, uso y préstamo en un ranking reproducible, medir sus resultados cuando las etiquetas maduren y mantener evidencia de cada versión del modelo.")
    add_heading(doc, "8 4 Estrategias", 2)
    add_bullets(doc, [
        "Certificar la madurez del TARGET antes de incorporar un periodo al entrenamiento.",
        "Seleccionar el top K diario según capacidad y evitar depender de un umbral fijo.",
        "Registrar contacto, resultado, queja y recuperación para medir impacto real.",
        "Vigilar drift, calibración y desempeño por periodo antes de mantener activo el modelo.",
    ])
    add_heading(doc, "8 5 Actividades", 2)
    add_table(doc, ["Actividad", "Producto verificable"], [
        ["Preparación de base", "Esquemas Bronze, Silver, Gold y API"],
        ["Ingesta", "bronze.telecom_raw y lote trazable"],
        ["Limpieza", "silver.telecom_operations"],
        ["Variables", "silver.behavior_features"],
        ["Contratos ML", "gold.ml_training y gold.ml_scoring"],
        ["Modelamiento", "Registro, métricas, importancia y predicciones"],
        ["Consumo", "Seis vistas API y dashboard Streamlit"],
    ], widths=[Cm(5), Cm(11)])

    add_heading(doc, "9 Método de marco lógico", 1)
    add_table(doc, ["Nivel", "Descripción", "Indicador", "Verificación"], [
        ["Fin", "Mejorar la asignación de capacidad de cobranza", "Riesgo capturado al 10 por ciento", "Métricas OOT"],
        ["Propósito", "Ordenar operaciones por probabilidad de mora", "ROC AUC OOT", "Registro del modelo"],
        ["Componentes", "Capas, modelo y dashboard", "Conteos y vistas disponibles", "Supabase y auditoría"],
        ["Actividades", "Cargar, limpiar, modelar y publicar", "Ejecución 00 a 07", "Notebooks y scripts"],
    ], widths=[Cm(2.2), Cm(5.5), Cm(4.2), Cm(4.1)], font_size=8.5)

    add_heading(doc, "10 Resultados esperados y logrados", 1)
    add_table(doc, ["Resultado", "Criterio", "Estado V2"], [
        ["Arquitectura Medallion", "Capas creadas y pobladas", "Logrado"],
        ["Conciliación", "Mismo volumen en Bronze Silver y Gold", "209,593 filas"],
        ["Contrato de entrenamiento", "TARGET maduro y roles temporales", "150,768 filas"],
        ["Población de scoring", "Sin TARGET maduro", "58,825 filas"],
        ["Modelo activo", "Registro y métricas reproducibles", "Random Forest"],
        ["Validación OOT", "ROC AUC superior a baseline", "0.8435"],
        ["Priorización", "Captura al 10 por ciento", "33.67 por ciento"],
        ["Dashboard", "Seis hojas basadas en Gold API", "Operativo"],
    ], widths=[Cm(5.2), Cm(6.6), Cm(4.2)], font_size=8.8)

    add_heading(doc, "11 Investigación y análisis metodológico", 1)
    add_body(doc, "La metodología combina CRISP DM para el ciclo analítico y Medallion para la persistencia. La V2 evita duplicar reglas: SQL transforma, Python orquesta y modela, los notebooks explican y el dashboard presenta. Esta división reduce el esfuerzo de mantenimiento y permite localizar cada regla en un único archivo.")
    add_heading(doc, "11 1 Arquitectura Medallion", 2)
    add_table(doc, ["Capa", "Responsabilidad", "Producto principal"], [
        ["Bronze", "Copia fiel y trazabilidad", "telecom_raw"],
        ["Silver", "Tipos, calidad y variables", "telecom_operations y behavior_features"],
        ["Gold", "Contratos certificados", "ml_training, ml_scoring y resultados"],
        ["API", "Consumo controlado", "Vistas por hoja del dashboard"],
    ], widths=[Cm(2.8), Cm(6.2), Cm(7)])
    add_heading(doc, "11 2 Secuencia técnica", 2)
    add_table(doc, ["Orden", "Notebook", "Resultado"], [
        ["00", "Configuración y conexión", "Entorno y Supabase validados"],
        ["01", "Ingesta Bronze", "Lote fuente registrado"],
        ["02", "Limpieza Silver", "Datos tipados y auditados"],
        ["03", "Construcción Gold", "Contratos ML publicados"],
        ["04", "EDA y KPIs", "Hallazgos de negocio"],
        ["05", "Modelo predictivo", "Candidato evaluado"],
        ["06", "Publicación de scoring", "Modelo y scores activos"],
        ["07", "Validación del dashboard", "KPIs conciliados"],
    ], widths=[Cm(2), Cm(7), Cm(7)], font_size=8.8)
    add_heading(doc, "11 3 División temporal", 2)
    add_table(doc, ["Partición", "Periodo", "Filas", "Uso"], [
        ["TRAIN", "01 jun a 07 jul", "104,357", "Ajuste final"],
        ["EMBARGO", "08 jul a 12 jul", "14,809", "Separación por horizonte"],
        ["OOT", "13 jul a 23 jul", "31,602", "Evaluación futura"],
        ["SCORING", "24 jul a 21 ago", "58,825", "Cola sin TARGET maduro"],
    ], widths=[Cm(3), Cm(4.2), Cm(3.2), Cm(5.6)])
    add_body(doc, "Los cinco folds usan ventanas expansivas. Cada validación ocurre después del entrenamiento y conserva una separación coherente con el horizonte de cinco días. Imputación, clipping y escalamiento se ajustan dentro del fold correspondiente.")

    add_heading(doc, "12 Desarrollo de analítica descriptiva y predictiva", 1)
    add_heading(doc, "12 1 Analítica descriptiva", 2)
    add_body(doc, "La ventana madura presenta una tasa de mora aproximada de 17.35 por ciento, pero la prevalencia OOT asciende a 20.7 por ciento. La diferencia confirma que el riesgo cambia en el tiempo y que una partición aleatoria habría ocultado parte de esa variación. Las variables monetarias y de frecuencia tienen distribuciones asimétricas, por lo que el pipeline aplica transformaciones robustas dentro de cada fold.")
    add_body(doc, "Los perfiles con mora muestran menor actividad de recarga y señales de mayor recencia. Estos patrones sirven para priorizar, pero no permiten concluir que aumentar una recarga cause una reducción de la mora.")
    add_heading(doc, "12 2 Comparación de modelos", 2)
    add_picture(doc, figures["models"], "Figura 1 Comparación de ROC AUC medio en validación temporal")
    add_body(doc, "Random Forest obtuvo un ROC AUC medio de 0.8670, una desviación de 0.0246 y un peor fold de 0.8407. HistGradientBoosting alcanzó un promedio cercano, 0.8650, con mayor variabilidad. La regresión logística registró 0.8179 y se deterioró en los últimos periodos. La regla de selección privilegió el peor fold y la estabilidad cuando la diferencia promedio quedó dentro de 0.005.")
    add_heading(doc, "12 3 Evaluación fuera de tiempo", 2)
    add_table(doc, ["Métrica OOT", "Resultado", "Interpretación"], [
        ["ROC AUC", "0.8435", "Buena capacidad de ordenamiento"],
        ["Average Precision", "0.6010", "Supera ampliamente la prevalencia base"],
        ["Brier Score", "0.1344", "Probabilidades útiles con calibración perfectible"],
        ["Recall al umbral", "73.20 por ciento", "Captura alta con mayor volumen activado"],
        ["Precisión al umbral", "48.81 por ciento", "Casi uno de cada dos activados presenta mora"],
    ], widths=[Cm(4.5), Cm(3.3), Cm(8.2)], font_size=8.8)
    add_body(doc, "El umbral OOF de 0.4194 activa una proporción mayor en OOT debido al cambio de prevalencia y distribución. Por ello, la política operativa usa ranking por capacidad y conserva el umbral como evidencia de evaluación, no como cuota fija de contactos.")
    add_heading(doc, "12 4 Resultados por capacidad", 2)
    add_picture(doc, figures["capacity"], "Figura 2 Precisión y riesgo capturado según capacidad OOT")
    add_table(doc, ["Capacidad", "Contactos", "Morosos capturados", "Precisión", "Recall", "Lift"], [
        ["5 por ciento", "1,581", "1,229", "77.74 por ciento", "18.82 por ciento", "3.76x"],
        ["10 por ciento", "3,161", "2,198", "69.54 por ciento", "33.67 por ciento", "3.37x"],
        ["20 por ciento", "6,321", "3,716", "58.79 por ciento", "56.92 por ciento", "2.85x"],
    ], widths=[Cm(2.7), Cm(2.5), Cm(3.7), Cm(2.5), Cm(2.5), Cm(2.1)], font_size=8.3)
    add_heading(doc, "12 5 Interpretación del modelo", 2)
    add_picture(doc, figures["features"], "Figura 3 Importancia por permutación en Test OOT")
    add_body(doc, "La señal se concentra en monto y conteo de recargas, recencia de la última recarga y saldo previo. La importancia por permutación mide cuánto pierde el modelo al alterar cada variable dentro del periodo OOT. No demuestra causalidad ni debe usarse para justificar decisiones automáticas sobre una persona.")

    add_heading(doc, "13 Dashboard interactivo", 1)
    add_body(doc, "El dashboard utiliza Streamlit y Plotly con una navegación semejante a un tablero de inteligencia de negocios. Cada hoja responde una pregunta concreta y consume una vista API derivada de Gold. La interfaz no contiene reglas de limpieza ni recrea variables del modelo.")
    add_table(doc, ["Hoja", "Pregunta", "Consumible"], [
        ["Resumen ejecutivo", "Cuál es la situación general", "page_executive_summary"],
        ["Riesgo y mora", "Cuándo se concentra el problema", "page_risk_and_delinquency"],
        ["Comportamiento", "Qué señales distinguen perfiles", "page_customer_behavior"],
        ["Modelo predictivo", "Qué tan estable es el score", "page_predictive_model"],
        ["Priorización", "A quién contactar primero", "page_collection_priority"],
        ["Calidad", "Los consumibles cumplen controles", "page_data_quality"],
    ], widths=[Cm(4), Cm(6.2), Cm(5.8)], font_size=8.7)
    add_body(doc, "La hoja de priorización muestra solo SCORING y permite filtrar bandas de riesgo, definir capacidad y descargar la cola. Las métricas de validación continúan disponibles en Gold para auditoría, pero no se mezclan con la lista operativa.")

    add_heading(doc, "14 Gobierno monitoreo y continuidad", 1)
    add_bullets(doc, [
        "Mantener una sola versión ACTIVE en gold.model_registry.",
        "Monitorear ROC AUC, Average Precision, Brier Score y recall a capacidad cuando maduren nuevas etiquetas.",
        "Comparar distribución de scores y variables con la ventana de entrenamiento.",
        "Suspender el uso ante una nueva ruptura del TARGET o una caída material de desempeño.",
        "Registrar el resultado de cada contacto para estimar recuperación, costo y fricción.",
        "Revisar manualmente los casos antes de aplicar una acción sobre el cliente.",
    ])
    add_table(doc, ["Riesgo", "Control propuesto"], [
        ["TARGET incompleto", "Cuarentena y certificación de madurez"],
        ["Drift", "Comparación temporal de scores y variables"],
        ["Sobrecontacto", "Identificador de cliente y regla de frecuencia"],
        ["Falsos positivos", "Revisión humana y registro de resultado"],
        ["Uso indebido", "Política que limite el score a priorización"],
        ["Secreto expuesto", "Entrada oculta y rotación de credenciales"],
    ], widths=[Cm(5), Cm(11)])

    add_heading(doc, "Conclusiones", 1)
    add_body(doc, "EDRSS V2 demuestra que una arquitectura ordenada puede sostener el ciclo completo de datos, modelo y consumo. Bronze, Silver y Gold conciliaron 209,593 operaciones. El modelo activo conserva poder discriminativo fuera de tiempo y transforma ese resultado en una regla de capacidad comprensible para cobranza.")
    add_body(doc, "El principal resultado operativo es la concentración del riesgo. Con una capacidad del 10 por ciento, el ranking captura 33.67 por ciento de los fallos y multiplica por 3.37 la precisión esperada frente a una selección aleatoria. Esta evidencia justifica un piloto controlado, no una automatización punitiva.")
    add_body(doc, "La próxima mejora debe incorporar una llave de cliente, costos de contacto, recuperación monetaria y etiquetas maduras posteriores. Esos datos permitirán medir retorno, evitar contactos repetidos y recalibrar el modelo con evidencia reciente.")

    add_heading(doc, "Recomendaciones", 1)
    add_bullets(doc, [
        "Iniciar un piloto con top K diario y revisión humana.",
        "Certificar periódicamente la integridad y madurez del TARGET.",
        "Incorporar un identificador anonimizado para controlar recurrencia.",
        "Medir costo, recuperación, quejas y resultado por contacto.",
        "Recalibrar y reevaluar el modelo antes de ampliar la cobertura.",
    ])

    add_heading(doc, "Referencias utilizadas", 1)
    references = [
        "[1] Björkegren D y Grissen D 2020 Behavior Revealed in Mobile Phone Usage Predicts Credit Repayment The World Bank Economic Review 34 3 618 634 https://doi.org/10.1093/wber/lhz006",
        "[2] Berg T Burg V Gombović A y Puri M 2020 On the Rise of FinTechs Credit Scoring Using Digital Footprints The Review of Financial Studies 33 7 2845 2897 https://doi.org/10.1093/rfs/hhz099",
        "[3] Lessmann S Baesens B Seow H V y Thomas L C 2015 Benchmarking State of the Art Classification Algorithms for Credit Scoring European Journal of Operational Research 247 1 124 136 https://doi.org/10.1016/j.ejor.2015.05.030",
        "[4] Schröer C Kruse F y Gómez J M 2021 A Systematic Literature Review on Applying CRISP DM Process Model Procedia Computer Science 181 526 534 https://doi.org/10.1016/j.procs.2021.01.199",
        "[5] Breiman L 2001 Random Forests Machine Learning 45 5 32 https://doi.org/10.1023/A:1010933404324",
        "[6] Databricks Medallion Architecture https://www.databricks.com/glossary/medallion-architecture",
        "[7] OpenML Delinquency Telecom Dataset 43745 https://www.openml.org/d/43745",
        "[8] Supabase Database documentation https://supabase.com/docs/guides/database",
        "[9] Streamlit documentation https://docs.streamlit.io",
    ]
    for reference in references:
        paragraph = doc.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        paragraph.paragraph_format.space_after = Pt(5)

    add_heading(doc, "Anexos", 1)
    add_heading(doc, "Anexo A Inventario de entregables", 2)
    add_table(doc, ["Elemento", "Ubicación"], [
        ["SQL", "sql 00_setup a 06_quality"],
        ["Notebooks", "notebooks 00 a 07"],
        ["Código Python", "src y scripts"],
        ["Dashboard", "dashboard app y pages"],
        ["Modelo", "models RandomForest edrss v2"],
        ["Documentación", "docs arquitectura storytelling y estado"],
    ], widths=[Cm(5), Cm(11)])
    add_heading(doc, "Anexo B Variables excluidas", 2)
    add_table(doc, ["Variable", "Motivo"], [
        ["label", "Etiqueta fuente"],
        ["riesgo_mora", "TARGET de negocio"],
        ["pdate", "Control temporal no predictor"],
        ["pcircle", "Constante en la fuente"],
        ["payback30 y payback90", "Linaje insuficiente y posible leakage"],
        ["IDs y versiones", "Gobierno y trazabilidad"],
    ], widths=[Cm(5), Cm(11)])
    doc.add_page_break()
    add_heading(doc, "Anexo C Glosario", 2)
    add_table(doc, ["Término", "Definición"], [
        ["OOF", "Predicción obtenida fuera del fold de entrenamiento"],
        ["OOT", "Evaluación en un periodo futuro no usado para seleccionar el modelo"],
        ["Recall a K", "Proporción de morosos capturados dentro de una capacidad"],
        ["Lift", "Precisión del ranking dividida por la prevalencia"],
        ["Drift", "Cambio de distribución entre periodos"],
        ["Cuarentena", "Datos que no cumplen todavía condiciones para uso supervisado"],
    ], widths=[Cm(4), Cm(12)])

    # Word actualizará índice y campos al abrir.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = "Early Delinquency Risk Scoring System EDRSS"
    doc.core_properties.subject = "Informe del Proyecto Productivo V2"
    doc.core_properties.author = "Equipo del Proyecto Productivo"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
